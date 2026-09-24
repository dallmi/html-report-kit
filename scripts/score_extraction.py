"""Regression test for the html-dashboard-to-excel skill.

Takes the script from skills/html-dashboard-to-excel/SKILL.md (the version that ships), runs it on
every test dashboard and scores the workbooks against test-dashboards/expected.json. A dataset counts
as found when a sheet has the expected number of data rows and every expected numeric column sum
appears among that sheet's column sums. Workbooks go to a temporary folder, never into the repo.

    python3 scripts/build_test_dashboards.py   # once, or after changing the test dashboards
    python3 scripts/score_extraction.py
"""
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "skills" / "html-dashboard-to-excel" / "SKILL.md"
EXP = json.loads((ROOT / "test-dashboards" / "expected.json").read_text())

tmp = Path(tempfile.mkdtemp())
script = tmp / "html_dashboard_to_xlsx.py"
script.write_text(re.search(r"```python\n(.*)```\n\Z", SKILL.read_text(), re.S).group(1))
total = found = 0
for rel, datasets in EXP.items():
    src = (ROOT / "test-dashboards" / rel).resolve()
    out = tmp / f"{src.parent.name}-{src.stem}.xlsx"
    r = subprocess.run([sys.executable, str(script), str(src), "--out", str(out)], capture_output=True, text=True)
    sheets = {}
    if r.returncode == 0:
        for ws in load_workbook(out, read_only=True).worksheets[1:]:
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            rows = rows[:next((i for i, x in enumerate(rows) if all(v is None for v in x)), len(rows))]
            sums = [sum(v for v in col if isinstance(v, (int, float)) and not isinstance(v, bool))
                    for col in zip(*rows)] if rows else []
            sheets[ws.title] = (len(rows), sums)
    print(f"\n{rel}  [{'ok' if r.returncode == 0 else 'EXIT %d' % r.returncode}]")
    if r.returncode:
        print("  " + (r.stdout + r.stderr).strip()[-400:])
    for d in datasets:
        total += 1
        hit = next((name for name, (n, sums) in sheets.items() if n == d["rows"] and
                    all(any(abs(s - v) <= 1e-6 * max(1, abs(v)) for s in sums) for v in d["sums"].values())), None)
        found += bool(hit)
        print(f"  {'✓' if hit else '✗'} {d['dataset']:<52} {d['rows']:>5} rows → {hit or '—'}")
print(f"\nFOUND {found} of {total} expected datasets")
sys.exit(0 if found == total else 1)
