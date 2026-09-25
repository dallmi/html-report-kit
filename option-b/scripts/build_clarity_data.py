"""Build the dashboard data files: data/manifest.json + data/*.json.

This job replaces the Excel control panel. Inclusion, date conversion and
derived metrics happen here, once, so a data refresh never touches the page:
drop the new JSON into data/ and reload.

Sources
  --from-html FILE     an existing single-file dashboard (`const DATA = {...}`).
                       The migration path: the output is checked against that
                       DATA and the build fails on any difference.
  --demo               the synthetic demo dashboard in this repository; short for
                       --from-html DEMO_HTML --reach-ref-month DEMO_REF_MONTH.
  --bridge FILE        the bridge workbook. Not implemented here: the sheet
                       layout lives in the internal skeleton; see read_bridge().
  --year YYYY          reporting year when DATA has none (the original dashboard
                       hardcodes it in the page); default: inferred from the
                       headcount months and article dates.
  --packs FILE         pack list workbook, sheet 07-packs (optional).
  --inspect            show the structure of DATA against what this build expects
                       (names, types, counts; no values) and write nothing.
  --overrides FILE     manual rows the exports miss (default: overrides.yaml).

    python3 scripts/build_clarity_data.py --demo

Standard library only; openpyxl for --packs, PyYAML once overrides.yaml has entries.
"""
import argparse
import datetime as dt
import json
import math
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCHEMA_VERSION = "1.0"
# --demo: the synthetic demo dashboard and the reach reference month its figures are built with
DEMO_HTML = ROOT.parent / "dashboard" / "comms-intelligence-dashboard-v3-demo.html"
DEMO_REF_MONTH = "2026-06"

MAIL_DIMS = ["cc", "div", "map", "team", "ct", "tm", "cl", "pk", "tcl", "tid"]
MAIL_VALUES = ["t", "es", "op", "uc", "ms", "po", "pc", "pu", "th", "tp"]
PAGE_ATTRS = ["site", "url", "div", "ct", "ov", "th", "tp", "org", "rg"]
ARTICLE_FIELDS = ["y", "m", "date", "t", "au", "ov", "th", "tp", "ts", "rg", "ch", "pk", "tcl", "tid",
                  "v", "uv", "li", "co", "reach"]
VIDEO_FIELDS = ["y", "m", "t", "div", "lang", "v", "uv", "c1", "c25", "c50", "c75", "c100", "eng"]
# pack list columns (sheet 07-packs) -> JSON field; see docs in README
PACK_COLUMNS = {"pack id": "pack_id", "cluster prefix": "cluster_prefix", "pack no.": "pack_no",
                "pack": "pack", "cluster": "cluster", "lead": "lead", "lead team": "lead_team",
                "partner team": "partner_team", "divisions": "divisions", "regions": "regions",
                "objective": "objectives", "start": "start", "end": "end", "launch": "launch",
                "in_report": "in_report"}
PACK_LISTS = {"divisions", "regions", "objectives"}
PACK_DATES = {"start", "end", "launch"}
INTRANET_VIEW_FACTOR = 0.8  # intranet reach = page views x 0.8 / dashboard headcount


class BuildError(Exception):
    pass


# ---------------------------------------------------------------- dates
def excel_serial_to_iso(x):
    """Excel serial day number (1900 system) -> ISO date."""
    return (dt.date(1899, 12, 30) + dt.timedelta(days=int(math.floor(float(x))))).isoformat()


def to_iso(x):
    if x is None or x == "":
        return None
    if isinstance(x, (dt.datetime, dt.date)):
        return (x.date() if isinstance(x, dt.datetime) else x).isoformat()
    if isinstance(x, (int, float)):
        return excel_serial_to_iso(x)
    s = str(x).strip()
    for f in ("%Y-%m-%d", "%d %b %Y", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    raise BuildError(f"unrecognised date {x!r}")


# ---------------------------------------------------------------- sources
def read_inline_data(path):
    html = Path(path).read_text(encoding="utf-8")
    key = "const DATA = "
    i = html.find(key)
    if i < 0:
        raise BuildError(f"{path}: no `const DATA = ` block")
    if html[i + len(key):].lstrip().startswith("/*__DATA__*/null"):
        raise BuildError(f"{path}: DATA is empty (template without data?)")
    try:
        data, _ = json.JSONDecoder().raw_decode(html, i + len(key))
    except json.JSONDecodeError as e:
        raise BuildError(f"{path}: DATA block is not JSON ({e})")
    if not isinstance(data, dict) or "mailP" not in data:
        raise BuildError(f"{path}: DATA is empty (template without data?)")
    return data


def mail_columns(mp):
    """(unknown, missing) columns of a legacy mailP block. Unknown columns are not carried over."""
    have = set(mp) - {"dims", "cols"}
    want = {"q", "m"} | set(MAIL_VALUES)
    have_dims = set(mp.get("dims") or {})
    return (sorted(have - want) + sorted(f"dims.{k}" for k in have_dims - set(MAIL_DIMS)),
            sorted(want - have) + sorted(f"dims.{k}" for k in set(MAIL_DIMS) - have_dims))


# what each source dataset must have, and what it may have; everything else is carried over unchanged
SOURCE_FIELDS = {
    "arts": ([k for k in ARTICLE_FIELDS if k not in ("y", "date", "reach")] + ["ds", "q"], ["y"]),
    "vids": ([k for k in VIDEO_FIELDS if k != "y"] + ["q"], ["y"]),
    "pgLevel": (PAGE_ATTRS + ["v"], ["uv", "vis", "uvy", "vy", "li", "co"]),
    "pgDiv": (["url", "sp", "uv"], PAGE_ATTRS),
    "pgReg": (["url", "sp", "uv"], PAGE_ATTRS),
    "pages": ([], []),
    "links": ([], []),
}
TOP_KEYS = ["mailP", "arts", "vids", "pgLevel", "pgDiv", "pgReg", "pages", "links", "hcGeduld", "hcDash", "src"]
TOP_OPTIONAL = ["year", "asOf", "prior", "source"]


def kind(v):
    return ("null" if v is None else "bool" if isinstance(v, bool) else "num" if isinstance(v, (int, float))
            else "text" if isinstance(v, str) else "list" if isinstance(v, list) else "object")


def inspect(D, path):
    """Structure of DATA against what the build expects: names, types and counts, never values."""
    def line(label, text):
        print(f"  {label:<9}{text}")

    print(f"DATA in {Path(path).name}")
    line("keys", " ".join(D))
    miss = [k for k in TOP_KEYS if k not in D]
    unknown = [k for k in D if k not in TOP_KEYS + TOP_OPTIONAL]
    line("missing", " ".join(miss) or "-")
    if unknown:
        line("unknown", " ".join(f"{k}:{kind(D[k])}" for k in unknown) + "  (not used)")
    line("year", "present" if isinstance(D.get("year"), int) else f"not in DATA, inferred {infer_year(D)}")
    for k in ("asOf", "prior", "source"):
        line(k, "present" if D.get(k) else "-")
    mp = D.get("mailP")
    if isinstance(mp, dict):
        u, m = mail_columns(mp)
        print(f"\nmailP  {len(mp.get('es') or [])} rows")
        line("columns", " ".join(k for k in mp if k not in ("dims", "cols")))
        line("dims", " ".join(f"{k}({len(v)})" for k, v in (mp.get("dims") or {}).items()))
        line("unknown", (" ".join(u) + "  (NOT carried over)") if u else "-")
        line("missing", " ".join(m) or "-")
    for name, (required, optional) in SOURCE_FIELDS.items():
        rows = D.get(name)
        if not isinstance(rows, list):
            continue
        fields = {}
        for r in rows:
            for k, v in r.items():
                fields.setdefault(k, set()).add(kind(v))
        print(f"\n{name}  {len(rows)} rows")
        line("fields", " ".join(fields))
        if required or optional:
            u = [f"{k}:{'/'.join(sorted(fields[k]))}" for k in fields if k not in required + optional]
            line("unknown", (" ".join(u) + "  (carried over)") if u else "-")
            line("missing", " ".join(k for k in required if k not in fields) or "-")
    for name in ("hcGeduld", "hcDash"):
        hc = D.get(name) or {}
        months = sorted(hc)
        print(f"\n{name}  {len(months)} months" + (f"  {months[0]} .. {months[-1]}" if months else ""))
    if isinstance(D.get("src"), dict):
        print(f"\nsrc  {' '.join(D['src'])}")


def read_bridge(path):
    """Bridge workbook -> the same legacy DATA dict read_inline_data returns.

    Not implemented offline: the sheet and column layout of the bridge workbook
    lives in the internal build skeleton. Port it here and apply, per sheet,
    the `Include` flag before returning rows (ship only included rows).
    """
    raise BuildError("--bridge is not implemented in this copy; use --from-html, "
                     "or port the internal skeleton into read_bridge()")


def read_packs(path):
    try:
        import openpyxl
    except ImportError:
        raise BuildError("--packs needs openpyxl (pip install openpyxl)")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "07-packs" not in wb.sheetnames:
        raise BuildError(f"{path}: no sheet 07-packs (found {', '.join(wb.sheetnames)})")
    rows = wb["07-packs"].iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(rows)]
    norm = [h.lower() for h in head]
    missing = [c for c in PACK_COLUMNS if c not in norm]
    if missing:
        raise BuildError(f"{path}: sheet 07-packs lacks columns: {', '.join(missing)}")
    act = [(i, h[len("activities_"):]) for i, h in enumerate(norm) if h.startswith("activities_")]
    out = []
    for r in rows:
        if all(v is None for v in r):
            continue
        rec = {"activities": {}}
        for i, h in enumerate(norm):
            f = PACK_COLUMNS.get(h)
            if not f:
                continue
            v = r[i] if i < len(r) else None
            if f in PACK_LISTS:
                v = [s.strip() for s in str(v or "").replace(";", ",").split(",") if s.strip()]
            elif f in PACK_DATES:
                v = to_iso(v)
            elif f == "in_report":
                v = str(v).strip().lower() in ("1", "true", "yes", "y", "x")
            rec[f] = v
        for i, k in act:
            rec["activities"][k] = r[i] if i < len(r) else None
        if not rec.get("pack_id"):
            raise BuildError(f"{path}: a pack row has no Pack ID")
        if rec["in_report"]:  # inclusion happens here, once
            del rec["in_report"]
            out.append(rec)
    return out


def read_overrides(path):
    p = Path(path)
    if not p.exists():
        return {}
    text = p.read_text(encoding="utf-8")
    if not any(l.strip() and not l.lstrip().startswith("#") for l in text.splitlines()):
        return {}
    try:
        import yaml
    except ImportError:
        raise BuildError(f"{p.name} has entries; reading it needs PyYAML (pip install pyyaml)")
    data = yaml.safe_load(text) or {}
    unknown = set(data) - {"articles", "videos"}
    if unknown:
        raise BuildError(f"{p.name}: only articles and videos take overrides, not {', '.join(sorted(unknown))}")
    return data


# ---------------------------------------------------------------- transforms
def columnar(mail_by_year):
    """Merge the years' dictionary-encoded mailings into one dims/cols block with a year column."""
    dims = {k: [] for k in MAIL_DIMS}
    cols = {k: [] for k in ["y", "m"] + MAIL_VALUES + MAIL_DIMS}
    for y, mp in mail_by_year:
        n = len(mp["es"])
        for k in ["m"] + MAIL_VALUES:
            if len(mp[k]) != n:
                raise BuildError(f"mailings {y}: column {k} has {len(mp[k])} rows, expected {n}")
        for i in range(n):
            if mp["q"][i] != math.ceil(mp["m"][i] / 3):
                raise BuildError(f"mailings {y} row {i}: quarter {mp['q'][i]} does not match month {mp['m'][i]}")
        cols["y"] += [y] * n
        cols["m"] += mp["m"]
        for k in MAIL_VALUES:
            cols[k] += mp[k]
        for k in MAIL_DIMS:
            src = mp["dims"][k]
            remap = []
            for v in src:
                if v not in dims[k]:
                    dims[k].append(v)
                remap.append(dims[k].index(v))
            cols[k] += [remap[j] for j in mp["cols"][k]]
    return {"dims": dims, "cols": cols}


def extra(row, known):
    """Fields this build does not know: carried over unchanged, after the known ones."""
    return {k: v for k, v in row.items() if k not in known}


def check_year(kind, r, y):
    """Some exports carry the year on each row; it must be the year of the list the row is in."""
    if "y" in r and r["y"] != y:
        raise BuildError(f"{kind} {r.get('t')!r}: year {r['y']} in the {y} list")


def articles(rows, y, hc):
    out = []
    for a in rows:
        check_year("article", a, y)
        date = to_iso(a["ds"])
        if int(date[5:7]) != a["m"] or a["q"] != math.ceil(a["m"] / 3):
            raise BuildError(f"article {a['t']!r}: date {a['ds']} does not match month {a['m']} / quarter {a['q']}")
        h = hc.get(f"{y}-{a['m']:02d}")
        known = {**{k: a[k] for k in ARTICLE_FIELDS if k in a}, "y": y, "date": date,
                 "reach": a["uv"] / h if h else None}
        out.append({**{k: known[k] for k in ARTICLE_FIELDS}, **extra(a, ARTICLE_FIELDS + ["ds", "q"])})
    return out


def videos(rows, y):
    for v in rows:
        check_year("video", v, y)
        if v["q"] != math.ceil(v["m"] / 3):
            raise BuildError(f"video {v['t']!r}: quarter {v['q']} does not match month {v['m']}")
    return [{**{k: v[k] for k in VIDEO_FIELDS if k in v}, "y": y, **extra(v, VIDEO_FIELDS + ["q"])} for v in rows]


def split_rows(split, pages, name):
    """Visitor split rows repeat every page attribute; keep url + split + its figures and join in the front end."""
    by_url = {p["url"]: p for p in pages}
    out = []
    for r in split:
        p = by_url.get(r["url"])
        if p is None:
            raise BuildError(f"{name}: url {r['url']} is not in the page-level export")
        if any(k in r and r[k] != p[k] for k in PAGE_ATTRS):
            raise BuildError(f"{name}: attributes of {r['url']} differ from the page-level export")
        out.append({"url": r["url"], "sp": r["sp"], "uv": r["uv"], **extra(r, PAGE_ATTRS + ["sp", "uv"])})
    return out


def apply_overrides(ds, ov, fields):
    rows = (ov or {}).get("append") or []
    for r in rows:
        missing = [k for k in fields if k not in r and k != "reach"]
        if missing or not r.get("reason"):
            raise BuildError(f"override row {r.get('t')!r}: needs {', '.join(missing + ([] if r.get('reason') else ['reason']))}")
        r = {**r, "date": to_iso(r["date"])} if "date" in fields else r
        if "date" in fields and int(r["date"][5:7]) != r["m"]:
            raise BuildError(f"override row {r['t']!r}: date {r['date']} does not match month {r['m']}")
        ds.append({**{k: r.get(k) for k in fields}, "override": r["reason"]})
    return len(rows)


def infer_year(D):
    """The original (v2) dashboard hardcodes its year in the page, not in DATA: take the year
    most headcount months and article dates carry, as the v3 fill skill does."""
    years = {}
    for k in D.get("hcGeduld") or {}:
        if len(str(k)) == 7 and str(k)[:4].isdigit():
            years[int(str(k)[:4])] = years.get(int(str(k)[:4]), 0) + 1
    for a in D.get("arts") or []:
        try:
            y = int(to_iso(a.get("ds"))[:4])
        except (BuildError, TypeError):
            continue
        years[y] = years.get(y, 0) + 1
    return max(years, key=years.get) if years else None


def set_year(D, args):
    """DATA.year, else --year, else inferred. Returns where it came from, for the build log."""
    if isinstance(D.get("year"), int):
        if args.year and args.year != D["year"]:
            raise BuildError(f"--year {args.year} contradicts DATA.year {D['year']}")
        return None
    D["year"] = args.year or infer_year(D)
    if not D["year"]:
        raise BuildError("reporting year unknown: DATA has no year, and neither headcount months nor "
                         "article dates give one; pass --year YYYY")
    return "--year" if args.year else "inferred from headcount months and article dates"


def build(D, args):
    Y = D.get("year")
    if not isinstance(Y, int):
        raise BuildError("DATA.year missing")
    prior = D.get("prior") or None
    PY = prior["year"] if prior else None
    if prior and PY != Y - 1:
        raise BuildError(f"prior year {PY} is not {Y - 1}")
    hcG, hcD = D.get("hcGeduld") or {}, D.get("hcDash") or {}
    refs = sorted(k for k in hcD if hcG.get(k))
    ref = args.reach_ref_month or (refs[-1] if refs else None)
    if ref and ref not in refs:
        raise BuildError(f"--reach-ref-month {ref}: no headcount for that month ({', '.join(refs)})")

    mail = columnar([(Y, D["mailP"])] + ([(PY, prior["mailP"])] if prior else []))
    arts = articles(D["arts"], Y, hcG) + (articles(prior.get("arts") or [], PY, hcG) if prior else [])
    vids = videos(D["vids"], Y) + (videos(prior.get("vids") or [], PY) if prior else [])
    ov = read_overrides(args.overrides)
    n_ov = apply_overrides(arts, ov.get("articles"), ARTICLE_FIELDS) + apply_overrides(vids, ov.get("videos"), VIDEO_FIELDS)
    for a in arts:  # override rows get their reach here too, from the same formula
        if "override" in a:
            h = hcG.get(f"{a['y']}-{a['m']:02d}")
            a["reach"] = a["uv"] / h if h else None
    hc_ref = hcD.get(ref) if ref else None
    pages = [{**p, "reach": p["v"] * INTRANET_VIEW_FACTOR / hc_ref if hc_ref else None} for p in D["pgLevel"]]
    if len({p["url"] for p in pages}) != len(pages):
        raise BuildError("page-level export has duplicate URLs")
    pages_ds = {"pages": pages,
                "visitor_division": split_rows(D["pgDiv"], D["pgLevel"], "pgDiv"),
                "visitor_region": split_rows(D["pgReg"], D["pgLevel"], "pgReg")}
    clicks_ds = {"pages": D["pages"], "links": D["links"]}
    packs_ds = read_packs(args.packs) if args.packs else []

    def years(rows):
        c = {}
        for y in rows:
            c[str(y)] = c.get(str(y), 0) + 1
        return c

    cutoff = to_iso(D["asOf"]) if D.get("asOf") else None
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "generated_at": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z"),
        "reporting_year": Y,
        "reach_reference_month": ref,
        "cutoff_date": cutoff,
        "sources": {"bridge": args.source_label or D.get("source") or Path(args.from_html).name,
                    "packs": Path(args.packs).name if args.packs else None},
        "headcount": {"geduld": hcG, "dashboard": hcD},
        "datasets": {
            "mailings": {"file": "mailings.json", "rows": len(mail["cols"]["y"]), "years": years(mail["cols"]["y"])},
            "articles": {"file": "articles.json", "rows": len(arts), "years": years(a["y"] for a in arts)},
            "videos": {"file": "videos.json", "rows": len(vids), "years": years(v["y"] for v in vids)},
            "pages": {"file": "pages.json", "rows": len(pages),
                      "tables": {k: len(v) for k, v in pages_ds.items()}},
            "clicks": {"file": "clicks.json", "rows": len(D["links"]),
                       "tables": {k: len(v) for k, v in clicks_ds.items()}},
            "packs": {"file": "packs.json", "rows": len(packs_ds)},
        },
    }
    files = {"mailings.json": mail, "articles.json": arts, "videos.json": vids,
             "pages.json": pages_ds, "clicks.json": clicks_ds, "packs.json": packs_ds}
    return manifest, files, n_ov


# ---------------------------------------------------------------- verification
def rebuild_legacy(manifest, files):
    """Invert build(): the legacy DATA shape, from the written files (override rows left out)."""
    Y = manifest["reporting_year"]
    mail = files["mailings.json"]

    def records(y):
        d, c = mail["dims"], mail["cols"]
        out = []
        for i, yy in enumerate(c["y"]):
            if yy != y:
                continue
            r = {k: c[k][i] for k in ["m"] + MAIL_VALUES}
            r.update({k: d[k][c[k][i]] for k in MAIL_DIMS})
            r["q"] = math.ceil(r["m"] / 3)
            out.append(r)
        return out

    # y stays on the rebuilt rows: the comparison only looks at fields the source rows have
    def art(a):
        d = dt.date.fromisoformat(a["date"])
        r = {k: v for k, v in a.items() if k not in ("date", "reach")}
        return {**r, "ds": d.strftime("%d %b %Y"), "q": math.ceil(a["m"] / 3)}

    def vid(v):
        return {**v, "q": math.ceil(v["m"] / 3)}

    real = lambda rows: [r for r in rows if "override" not in r]
    pg = files["pages.json"]
    by_url = {p["url"]: {k: p[k] for k in PAGE_ATTRS} for p in pg["pages"]}
    out = {"mail": {y: records(y) for y in sorted(set(mail["cols"]["y"]))},
           "arts": {}, "vids": {},
           "pgLevel": [{k: v for k, v in p.items() if k != "reach"} for p in pg["pages"]],
           "pgDiv": [{**by_url[r["url"]], **r} for r in pg["visitor_division"]],
           "pgReg": [{**by_url[r["url"]], **r} for r in pg["visitor_region"]],
           "pages": files["clicks.json"]["pages"], "links": files["clicks.json"]["links"],
           "hcGeduld": manifest["headcount"]["geduld"], "hcDash": manifest["headcount"]["dashboard"],
           "year": Y, "asOf": manifest["cutoff_date"]}
    for a in real(files["articles.json"]):
        out["arts"].setdefault(a["y"], []).append(art(a))
    for v in real(files["videos.json"]):
        out["vids"].setdefault(v["y"], []).append(vid(v))
    return out


def legacy_records(mp):
    out = []
    for i in range(len(mp["es"])):
        r = {k: mp[k][i] for k in ["q", "m"] + MAIL_VALUES}
        r.update({k: mp["dims"][k][mp["cols"][k][i]] for k in MAIL_DIMS})
        out.append(r)
    return out


def canon(rows, order):
    """Legacy rows with keys in a fixed order, so dict equality reports the first real difference."""
    return [{k: r[k] for k in order if k in r} for r in rows]


def verify(D, manifest, files):
    got = rebuild_legacy(manifest, files)
    Y, prior = D["year"], D.get("prior")
    checks = [("mailings", legacy_records(D["mailP"]), got["mail"].get(Y, [])),
              ("articles", D["arts"], got["arts"].get(Y, [])),
              ("videos", D["vids"], got["vids"].get(Y, [])),
              ("pgLevel", D["pgLevel"], got["pgLevel"]), ("pgDiv", D["pgDiv"], got["pgDiv"]),
              ("pgReg", D["pgReg"], got["pgReg"]), ("pages", D["pages"], got["pages"]),
              ("links", D["links"], got["links"]),
              ("hcGeduld", D["hcGeduld"], got["hcGeduld"]), ("hcDash", D["hcDash"], got["hcDash"]),
              ("year", D["year"], got["year"]), ("asOf", D.get("asOf"), got["asOf"])]
    if prior:
        PY = prior["year"]
        checks += [("prior mailings", legacy_records(prior["mailP"]), got["mail"].get(PY, [])),
                   ("prior articles", prior.get("arts") or [], got["arts"].get(PY, [])),
                   ("prior videos", prior.get("vids") or [], got["vids"].get(PY, []))]
    src = D.get("src") or {}
    for k, n in [("mail", len(D["mailP"]["es"])), ("arts", len(D["arts"])), ("vids", len(D["vids"])),
                 ("pages", len(D["pages"])), ("links", len(D["links"])), ("pgLevel", len(D["pgLevel"]))]:
        if k in src and src[k] != n:
            raise BuildError(f"source DATA.src.{k} = {src[k]} but the dataset has {n} rows")
    bad = []
    for name, want, have in checks:
        if isinstance(want, list):
            keys = list(dict.fromkeys(k for r in want for k in r))
            want, have = canon(want, keys), canon(have, keys)
            if len(want) != len(have):
                bad.append(f"{name}: {len(have)} rows, source has {len(want)}")
                continue
            first = {}  # every field that differs anywhere, with its first row, so one run shows them all
            for i, (a, b) in enumerate(zip(want, have)):
                if a != b:
                    for k in keys:
                        if k not in first and a.get(k) != b.get(k):
                            first[k] = (i, a.get(k), b.get(k))
            for k, (i, x, y) in first.items():
                bad.append(f"{name} row {i}, field {k}: source {x!r}, built {y!r}")
        elif want != have:
            bad.append(f"{name}: source {want!r}, built {have!r}")
    if bad:
        raise BuildError("built data does not match the source DATA:\n  " + "\n  ".join(bad))
    return len(checks)


# ---------------------------------------------------------------- write
def write(out_dir, manifest, files):
    """Each file replaced atomically; manifest last, so it never announces files that are not there yet."""
    out_dir.mkdir(parents=True, exist_ok=True)

    def put(name, obj):
        fd, tmp = tempfile.mkstemp(dir=out_dir, prefix="." + name, suffix=".tmp")
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            json.dump(obj, fh, ensure_ascii=False, separators=(",", ":"))
        os.chmod(tmp, 0o644)  # served by a web server or read from a mounted volume
        os.replace(tmp, out_dir / name)

    for name, obj in files.items():
        put(name, obj)
    fd, tmp = tempfile.mkstemp(dir=out_dir, prefix=".manifest", suffix=".tmp")
    with os.fdopen(fd, "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    os.chmod(tmp, 0o644)
    os.replace(tmp, out_dir / "manifest.json")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--from-html", help="single-file dashboard with an inline DATA block")
    src.add_argument("--demo", action="store_true", help="the synthetic demo dashboard in this repository")
    src.add_argument("--bridge", help="bridge workbook (not implemented in this copy)")
    ap.add_argument("--packs", help="pack list workbook with sheet 07-packs")
    ap.add_argument("--overrides", default=str(ROOT / "overrides.yaml"))
    ap.add_argument("--year", type=int, help="reporting year, when DATA has none and it cannot be inferred")
    ap.add_argument("--reach-ref-month", help="YYYY-MM; default: latest month with both headcounts")
    ap.add_argument("--source-label", help="provenance line shown in the dashboard header")
    ap.add_argument("--out", default=str(ROOT / "data"))
    ap.add_argument("--inspect", action="store_true",
                    help="show the structure of DATA (names, types, counts; no values) and write nothing")
    args = ap.parse_args()
    if args.demo:
        args.from_html = str(DEMO_HTML)
        args.reach_ref_month = args.reach_ref_month or DEMO_REF_MONTH
    try:
        D = read_inline_data(args.from_html) if args.from_html else read_bridge(args.bridge)
        if args.inspect:
            return inspect(D, args.from_html or args.bridge)
        year_from = set_year(D, args)
        manifest, files, n_ov = build(D, args)
        n = verify(D, manifest, files)
        write(Path(args.out), manifest, files)
    except BuildError as e:
        sys.exit(f"build failed: {e}")
    if year_from:
        print(f"reporting year {manifest['reporting_year']} ({year_from})")
    for label, mp in [("mailings", D["mailP"])] + ([("prior mailings", D["prior"]["mailP"])] if D.get("prior") else []):
        unknown = mail_columns(mp)[0]
        if unknown:
            print(f"note: {label} columns not carried over: {' '.join(unknown)} (run --inspect)")
    print(f"verified {n} checks against the source DATA: identical")
    for k, v in manifest["datasets"].items():
        print(f"  {v['file']:<14} {v['rows']:>6} rows" + (f"  {v['years']}" if "years" in v else ""))
    if n_ov:
        print(f"  {n_ov} override rows appended from {Path(args.overrides).name}")
    print(f"wrote {Path(args.out)}")


if __name__ == "__main__":
    main()
