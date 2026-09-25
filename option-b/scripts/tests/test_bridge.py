"""--bridge reads the bridge workbook into exactly the DATA the demo dashboard carries.

    python3 -m unittest discover -s option-b/scripts/tests

The test writes a bridge workbook and a pack list from the synthetic demo DATA, laid out like the
real ones (same sheets, headers and unused columns), reads them back with read_bridge and compares
every record with the demo. Two things the demo cannot carry are generated instead: tracking IDs
that encode pack and cluster the way the real ones do, and one theme and topic per mailing.
Needs openpyxl.
"""
import argparse
import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import build_clarity_data as b  # noqa: E402

try:
    import openpyxl
except ImportError:
    openpyxl = None

MAIL_HEAD = ["Include", "Year", "Quarter", "Month", "Corp Comms", "Corp Comms team", "iMEP object structure",
             "projectname", "foldername", "Title", "Mailings sent", "Emails sent", "Emails sent per mailing",
             "Avg % unique opens", "Avg % unique clicks", "Avg % click-to-open rate", "Mailing Sent Date",
             "Templates", "Content Type", "Theme", "Topic", "Tracking ID", "#open", "#unique clicks",
             "Mailing cluster", "Corp Comms new mapping", "Corp Comms KPI reporting"]
ART_HEAD = ["Include", "quarter", "week of year", "article title", "article short url", "publishing date",
            "overall views", "UV_EN", "UV_DE", "UV_FR", "UV_IT", "overall likes", "overall comments", "engagement",
            "engagement rate", "abstract personalization", "top story", "overtitle", "theme tag", "topic tag",
            "org announcement", "author", "Tracking ID", "News channel", "UV_total", "_Month", "_Quarter"]
VID_HEAD = ["Include", "video name", "asset ID", "created date", "language", "duration", "unique visitors", "views",
            "total time viewed", "views at 1%", "views at 25%", "views at 50%", "views at 75%", "views at 100%",
            "engagement score", "business division owner", "_Month", "_Quarter"]
WINDOWS = ["", "7", "14", "21", "28", "ytd"]
PAGE_HEAD = (["Include", "Site name", "Business division owner", "Page URL", "Content Type Tag", "Overtitle", "Theme",
              "Topic Tag"] + [f"uniquevisitor{w}" for w in WINDOWS] + [f"visits{w}" for w in WINDOWS]
             + [f"views{w}" for w in WINDOWS] + ["likes", "comments", "Target organization", "Target region"])
CLICK_HEAD = ["Include", "site name", "page name", "page URL", "links", "views", "unique visitors", "clicks",
              "unique clicks", "downloads", "ctvr", "uctuvr"]
LINK_HEAD = ["Include", "page name", "page URL", "link title", "destination URL", "clicks", "unique clicks",
             "downloads", "ctvr", "uctuvr"]
PACK_HEAD = ["Pack ID", "Cluster prefix", "Pack no.", "Pack", "Cluster", "Lead", "Lead team", "Partner team",
             "Divisions", "Regions", "Objective"]  # no in_report column, as in the photographed list


class Tracking:
    """Tracking IDs for (pack, cluster) pairs, and the pack list that resolves them back."""

    def __init__(self):
        self.prefix, self.pack_id, self.rows, self.n = {}, {}, [], 0

    def pfx(self, tcl):
        if tcl not in self.prefix:
            self.prefix[tcl] = f"CL{len(self.prefix):03d}" if tcl else "NOCLU"
        return self.prefix[tcl]

    def tid(self, pk, tcl):
        self.n += 1
        if not pk and not tcl:
            return "CCCCC-0000000-%05d" % self.n if self.n % 5 == 0 else "", ""  # placeholder: no tracking ID
        p = self.pfx(tcl)
        if not pk:
            return f"{p}-{b.NO_PACK}-{self.n:05d}", None
        if (pk, tcl) not in self.pack_id:
            pid = f"{p}-{len(self.pack_id) + 1:07d}"
            self.pack_id[pk, tcl] = pid
            first = not any(r[1] == p for r in self.rows)  # later packs of a prefix leave Cluster blank
            self.rows.append([pid, p, len(self.pack_id), pk, tcl if first else None] + [None] * 6)
        return f"{self.pack_id[pk, tcl]}-{self.n:05d}", None

    def pack_rows(self):
        rows = list(self.rows)
        for tcl, p in self.prefix.items():  # a cluster used only without a pack still needs its name
            if tcl and not any(r[1] == p and r[4] for r in rows):
                rows.append([f"{p}-9999999", p, 9999999, "(cluster row)", tcl] + [None] * 6)
        rows.append(["CCCCC-0000475", "CCCCC", 475, "Placeholder-cluster pack", None] + [None] * 6)
        return rows


def expected_tid(generated):
    tid, _ = generated
    return "" if tid.startswith("CCCCC-0000000") else tid


def write_fixture(D, folder):
    """Bridge workbook + pack list from demo DATA. Returns their paths and the tracking IDs written per row."""
    tr, tids = Tracking(), {"mail": {}, "arts": {}}
    wb = openpyxl.Workbook()
    wb.active.title = "Read me"

    def sheet(name, head, rows):
        ws = wb.create_sheet(name)
        ws.append(head)
        for r in rows:
            ws.append(r)

    mail_rows = []
    for y, mp in [(D["year"], D["mailP"]), (D["prior"]["year"], D["prior"]["mailP"])]:
        for i, r in enumerate(b.legacy_records(mp)):
            gen = tr.tid(r["pk"], r["tcl"])
            tids["mail"][y, i] = expected_tid(gen)
            # the original dashboard's rule for a blank KPI reporting cell, spelled out rather than taken from b
            blank_div = {"Not Corp Comms": "Non Corp Comms", "Corp Comms": "Corp Comms – unmapped"}[r["cc"]]
            div = "" if r["div"] == blank_div else r["div"]
            mail_rows.append([1, y, r["q"], r["m"], r["cc"], r["team"], "struct", "project", "folder", r["t"], r["ms"],
                              r["es"], r["es"] / max(r["ms"], 1), r["po"], r["pu"], r["pc"], dt.datetime(y, r["m"], 3),
                              r["tm"], r["ct"], (r["th"] or [None])[0], (r["tp"] or [None])[0], gen[0] or None,
                              r["op"], r["uc"], r["cl"], "" if r["map"] == "Not mapped" else r["map"], div])
    sheet("Data_Mailings", MAIL_HEAD, mail_rows)

    art_rows = []
    for y, arts in [(D["year"], D["arts"]), (D["prior"]["year"], D["prior"]["arts"])]:
        for i, a in enumerate(arts):
            gen = tr.tid(a["pk"], a["tcl"])
            tids["arts"][y, i] = expected_tid(gen)
            date = dt.datetime.strptime(a["ds"], "%d %b %Y")
            email = ".".join(a["au"].lower().split(" ")) + "@example.com"
            art_rows.append([1, a["q"], 1, a["t"], "/short", date, a["v"], a["uv"], None, None, None, a["li"], a["co"],
                             a["li"] + a["co"], 0.01, a["rg"], a["ts"] or None, a["ov"], a["th"], ",".join(a["tp"]),
                             "No", email, gen[0] or None, a["ch"], a["uv"], a["m"], a["q"]])
    sheet("Data_Articles", ART_HEAD, art_rows)

    vid_rows = []
    for y, vids in [(D["year"], D["vids"]), (D["prior"]["year"], D["prior"]["vids"])]:
        for v in vids:
            vid_rows.append([1, v["t"], "ASSET-1", dt.datetime(y, v["m"], 15), v["lang"], "00:01:00", v["uv"], v["v"],
                             "00d", v["c1"], v["c25"], v["c50"], v["c75"], v["c100"], v["eng"], v["div"], v["m"], v["q"]])
    sheet("Data_Videos", VID_HEAD, vid_rows)

    def page_cells(p):
        return [p["site"], p["div"], p["url"], p["ct"], p["ov"], p["th"], ",".join(p["tp"])]

    sheet("Data_Pages_level", PAGE_HEAD,
          [[1] + page_cells(p) + [p["uv"], 1, 2, 3, 4, p["uvy"], p["vis"], 1, 2, 3, 4, 5, p["v"], 1, 2, 3, 4, p["vy"],
                                  p["li"], p["co"], p["org"], p["rg"]] for p in D["pgLevel"]])
    # visitor splits: no YTD column (the demo splits carry none), tag columns as the page has them
    split_head = PAGE_HEAD[:8] + ["{}"] + [f"uniquevisitor{w}" for w in WINDOWS[:-1]] + ["visits", "views"]
    for name, label, rows in [("Data_Pages_Divisional_Split", "Visitor Division", D["pgDiv"]),
                              ("Data_Pages_Regional_Split", "Visitor Region", D["pgReg"])]:
        sheet(name, [label if h == "{}" else h for h in split_head],
              [[1] + page_cells(r) + [r["sp"], r["uv"], 1, 2, 3, 4, 9, 9] for r in rows])
    sheet("Data_Clicks_Pages", CLICK_HEAD,
          [[1, p["site"], p["p"], p["url"], 3, p["v"], p["uv"], p["c"], p["ucl"], p["d"], 0.1, 0.1] for p in D["pages"]])
    sheet("Data_Clicks_Links", LINK_HEAD,
          [[1, l["p"], l["url"], l["t"], l["dest"], l["c"], l["ucl"], l["d"], 0.1, 0.1] for l in D["links"]])

    hc = wb.create_sheet("Ref_Headcount")
    hc.append([None, "Reference headcount"])
    hc.append([None, "Two separate reach denominators — do not merge"])
    hc.append([])
    hc.append([None, "Geduld headcount (articles reach)"])
    hc.append([None, "Key", "Year", "Month", "Status", "Count"])
    for i, (k, n) in enumerate(sorted(D["hcGeduld"].items())):
        y, m = int(k[:4]), int(k[5:])
        key = f"{y - 1}-{m:02d}" if i == 0 else k  # the January slip in the real sheet: wrong Key on All
        for status, count in [("All", n + 30000), ("Active", n - 2400), ("Paid leave", 1800),
                              ("Unpaid leave", 600), ("Internal", n)]:
            hc.append([None, key if status == "All" else k, y, m, status, count])
    hc.append([])
    hc.append([None, "Dashboard headcount (intranet reach)"])
    hc.append([None, "Key", "Year", "Month", "Count"])
    for k, n in sorted(D["hcDash"].items()):
        hc.append([None, k, int(k[:4]), int(k[5:]), n])
    wb.create_sheet("Calc")
    bridge = Path(folder) / "bridge.xlsx"
    wb.save(bridge)

    pw = openpyxl.Workbook()
    pw.active.title = b.PACK_SHEET
    pw.active.append(PACK_HEAD)
    for r in tr.pack_rows():
        pw.active.append(r)
    packs = Path(folder) / "packs.xlsx"
    pw.save(packs)
    return bridge, packs, tids


@unittest.skipIf(openpyxl is None, "needs openpyxl")
class BridgeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.demo = b.read_inline_data(b.DEMO_HTML)
        cls.bridge, cls.packs, cls.tids = write_fixture(cls.demo, cls.tmp.name)
        b.NOTES.clear()
        cls.D = b.read_bridge(cls.bridge, cls.packs, cutoff=cls.demo["asOf"])
        cls.notes = list(b.NOTES)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def assertRows(self, got, want, name):
        self.assertEqual(len(got), len(want), name)
        for i, (g, w) in enumerate(zip(got, want)):
            self.assertEqual(g, w, f"{name} row {i}")

    def test_mailings(self):
        for y, mp, got in [(self.demo["year"], self.demo["mailP"], self.D["mailP"]),
                           (self.demo["prior"]["year"], self.demo["prior"]["mailP"], self.D["prior"]["mailP"])]:
            want = [{**r, "th": r["th"][:1], "tp": r["tp"][:1], "tid": self.tids["mail"][y, i],
                     "pk": r["pk"] if self.tids["mail"][y, i] else "",
                     "tcl": r["tcl"] if self.tids["mail"][y, i] else ""}
                    for i, r in enumerate(b.legacy_records(mp))]
            self.assertRows(b.legacy_records(got), want, f"mailings {y}")

    def test_articles(self):
        for y, want_rows, got in [(self.demo["year"], self.demo["arts"], self.D["arts"]),
                                  (self.demo["prior"]["year"], self.demo["prior"]["arts"], self.D["prior"]["arts"])]:
            for i, (g, w) in enumerate(zip(got, want_rows)):
                tid = self.tids["arts"][y, i]
                self.assertEqual(g["aus"], w["au"], f"articles {y} row {i}: display name")
                self.assertTrue(g["au"].endswith("@example.com"))
                want = {**w, "au": g["au"], "aus": w["au"], "tid": tid, "pk": w["pk"] if tid else "",
                        "tcl": w["tcl"] if tid else ""}
                self.assertEqual(g, want, f"articles {y} row {i}")
            self.assertEqual(len(got), len(want_rows))

    def test_videos_pages_clicks(self):
        self.assertRows(self.D["vids"], self.demo["vids"], "videos")
        self.assertRows(self.D["prior"]["vids"], self.demo["prior"]["vids"], "prior videos")
        self.assertRows(self.D["pgLevel"], self.demo["pgLevel"], "pgLevel")
        for k in ("pgDiv", "pgReg"):
            self.assertRows(self.D[k], [{"url": r["url"], "uv": r["uv"], "sp": r["sp"]} for r in self.demo[k]], k)
        self.assertRows(self.D["pages"], self.demo["pages"], "click pages")
        self.assertRows(self.D["links"], self.demo["links"], "links")

    def test_top_level(self):
        for k in ("year", "asOf", "hcGeduld", "hcDash"):
            self.assertEqual(self.D[k], self.demo[k], k)
        self.assertEqual(self.D["prior"]["year"], self.demo["prior"]["year"])
        self.assertEqual(self.D["src"], self.demo["src"])

    def test_notes(self):
        text = "\n".join(self.notes)
        self.assertIn("Key does not match Year-Month in 1 row", text)
        self.assertNotIn("≠", text)  # internal = active + paid + unpaid in every month of the fixture
        self.assertNotIn("not in the pack list", text)

    def test_build_writes_the_demo_datasets(self):
        """The whole pipeline on the bridge DATA: build, round-trip check, same row counts as --demo."""
        args = argparse.Namespace(reach_ref_month=b.DEMO_REF_MONTH, overrides=str(Path(self.tmp.name) / "none.yaml"),
                                  packs=str(self.packs), source_label=None, from_html=None, bridge=str(self.bridge))
        D = b.read_bridge(self.bridge, self.packs, cutoff=self.demo["asOf"])
        manifest, files, _ = b.build(D, args)
        b.verify(D, manifest, files)
        demo_manifest, _, _ = b.build(self.demo, argparse.Namespace(**{**vars(args), "packs": None,
                                                                      "from_html": str(b.DEMO_HTML)}))
        rows = lambda m: {k: (v["rows"], v.get("years"), v.get("tables")) for k, v in m["datasets"].items() if k != "packs"}
        self.assertEqual(rows(manifest), rows(demo_manifest))
        self.assertEqual(manifest["sources"]["bridge"], "bridge.xlsx")
        self.assertEqual(manifest["datasets"]["packs"]["rows"], len(files["packs.json"]))
        self.assertTrue(all("pack_id" in p for p in files["packs.json"]))


class UnitTest(unittest.TestCase):
    def test_tracking(self):
        packs, clusters = {"3KEYS-0000444": "Sustainable Impact", "CCCCC-0000475": "Survey"}, {"3KEYS": "Three Keys"}
        unknown = set()
        t = lambda tid: b.tracking(tid, packs, clusters, unknown)
        self.assertEqual(t("3KEYS-0000444-001"), ("3KEYS-0000444-001", "Sustainable Impact", "Three Keys"))
        self.assertEqual(t("3KEYS-0000000-002"), ("3KEYS-0000000-002", "", "Three Keys"))
        self.assertEqual(t("CCCCC-0000000-003"), ("", "", ""))
        self.assertEqual(t("CCCCC-0000475-004"), ("CCCCC-0000475-004", "Survey", ""))
        self.assertEqual(t(""), ("", "", ""))
        self.assertEqual(t("3KEYS-0000999-005"), ("3KEYS-0000999-005", "", "Three Keys"))
        self.assertEqual(unknown, {"3KEYS-0000999"})

    def test_author_name(self):
        self.assertEqual(b.author_name("anna-lena.vogt@example.com"), "Anna-Lena Vogt")
        self.assertEqual(b.author_name("omar.haddad@example.com"), "Omar Haddad")
        self.assertEqual(b.author_name("Comms team"), "Comms team")

    def test_convert_blanks(self):
        self.assertEqual(b.convert(None, "count", "x"), 0)  # counts: blank is zero
        self.assertIsNone(b.convert(None, "rate", "x"))  # rates: blank is no data, left out of averages
        self.assertEqual(b.convert(584.0, "count", "x"), 584)
        self.assertEqual(b.convert(" GIC ", "text", "x"), "GIC")
        self.assertEqual(b.convert(None, "tag", "x"), [])
        self.assertEqual(b.convert("Agile,Artificial Intelligence", "tags", "x"), ["Agile", "Artificial Intelligence"])
        with self.assertRaises(b.BuildError):
            b.convert("n/a", "count", "x")


if __name__ == "__main__":
    unittest.main()
